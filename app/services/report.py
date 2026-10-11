from datetime import datetime
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Event
from app.services.catalog import active_cards, analyze_card, market_as_of


def build_report(db: Session) -> dict:
    as_of = market_as_of(db)
    cards = []
    for card in active_cards(db):
        metrics = analyze_card(db, card, end=as_of)
        cards.append(_card_chapter(metrics))
    ranked = sorted(cards, key=lambda item: item["advice"]["score"], reverse=True)
    events = [
        {
            "id": row.id,
            "date": row.event_date.isoformat(),
            "type": row.event_type,
            "card_code": row.card_code,
            "title": row.title,
            "description": row.description,
            "impact_level": row.impact_level,
            "source_url": row.source_url,
        }
        for row in db.scalars(select(Event).order_by(Event.event_date.desc())).all()
    ]
    return {
        "title": "集换社卡牌交易量化分析报告",
        "generated_at": datetime.now().isoformat(timespec="minutes"),
        "as_of": as_of.isoformat() if as_of else None,
        "disclaimer": "本报告只统计 PSA10 且非套出成交，价格单位为人民币。规则引擎输出供复盘，不构成投资建议。",
        "executive": _executive(ranked, events),
        "metrics_table": [_table_row(item) for item in ranked],
        "cards": ranked,
        "divergence": _divergence_table(ranked),
        "liquidity": _liquidity(ranked),
        "anomalies": [
            {"card_code": item["card_code"], "card_name": item["card_name"], "count": item["anomaly_count"], "rows": item["anomalies"][:12]}
            for item in ranked
            if item["anomaly_count"]
        ],
        "forecasts": [_forecast_row(item) for item in ranked],
        "advice": {
            "ranking": [_advice_row(item) for item in ranked],
            "risk_controls": [
                "单卡仓位不超过该卡建议上限，三款高供给皮卡丘不要当成同一笔低相关资产。",
                "价升量缩或 7 日涨幅超过 20% 且缩量时，不追已经走出来的价格。",
                "套出成交和 IQR 异常点只作观察，不拿来当最新价或目标价。",
                "出现追加生产、再抽选等高影响供给事件时，把目标价下修一档。",
            ],
        },
        "external": {
            "note": "PSA Population 为人工维护的参考值，竞价是录入或导入的快照，可能滞后于集换社页面。",
            "population": [
                {
                    "card_code": item["card_code"],
                    "card_name": item["card_name"],
                    "total": (item.get("population") or {}).get("total"),
                    "psa10": (item.get("population") or {}).get("psa10"),
                    "gem_rate": (item.get("population") or {}).get("gem_rate"),
                    "pressure": ((item.get("population") or {}).get("pressure") or {}).get("label"),
                }
                for item in ranked
            ],
            "events": events,
        },
        "monitoring": [_monitor(item) for item in ranked],
    }


def _card_chapter(metrics: dict) -> dict:
    advice = metrics["advice"]
    public = dict(metrics)
    public["narrative"] = _narrative(metrics, advice)
    return public


def _narrative(metrics: dict, advice: dict) -> str:
    if not metrics.get("has_psa10"):
        return f"{metrics['card_name']} 目前没有可计价的 PSA10 成交。"
    ret7 = (metrics["returns"].get("d7") or {}).get("return")
    ret30 = (metrics["returns"].get("d30") or {}).get("return")
    volume = metrics["volume"]
    risk = metrics["risk"]
    pieces = [
        f"{metrics['card_name']} 最新 PSA10 日中位价 {_cny(metrics['latest_price'])}（{metrics['latest_date']}）。",
        f"样本 {metrics['sample_start']} 至 {metrics['sample_end']}，共 {metrics['psa10_count']} 笔真实 PSA10 成交。",
        f"7 日收益 {_pct(ret7)}，30 日收益 {_pct(ret30)}，90 日收益 {_pct((metrics['returns'].get('d90') or {}).get('return'))}。",
        f"近 7 日成交 {volume.get('d7')} 笔，量能变化 {_pct(volume.get('change_7d'))}，信号为{metrics['divergence']['signal']}。",
        f"波动率 {_pct(risk.get('volatility'))}，最大回撤 {_pct(risk.get('max_drawdown'))}，当前价格分位 {_pct(metrics.get('percentile'))}。",
        f"规则建议：{advice['action']}。{advice['position']}。",
    ]
    if metrics.get("bundle_count"):
        pieces.append(f"另有 {metrics['bundle_count']} 笔套出记录已剔除，不进入价格和收益。")
    if metrics.get("warnings"):
        pieces.append("注意：" + "、".join(metrics["warnings"]) + "。")
    return "".join(pieces)


def _executive(cards: list[dict], events: list[dict]) -> list[str]:
    usable = [item for item in cards if item.get("has_psa10")]
    if not usable:
        return ["关注列表里还没有可计算的 PSA10 成交。"]
    lines = []
    by_7d = [item for item in usable if (item["returns"]["d7"] or {}).get("return") is not None]
    if by_7d:
        lead = max(by_7d, key=lambda item: item["returns"]["d7"]["return"])
        lag = min(by_7d, key=lambda item: item["returns"]["d7"]["return"])
        lines.append(f"7 日涨幅最大的是 {lead['card_name']}（{_pct(lead['returns']['d7']['return'])}），最弱的是 {lag['card_name']}（{_pct(lag['returns']['d7']['return'])}）。")
    shrink = [item for item in usable if item["divergence"]["signal"] == "价升量缩"]
    if shrink:
        lines.append("价升量缩：" + "、".join(item["card_name"] for item in shrink) + "。这几张的上涨缺少成交配合。")
    risky = [item for item in usable if (item["risk"].get("max_drawdown") or 0) <= -0.35]
    if risky:
        lines.append("回撤较深：" + "、".join(f"{item['card_name']} {_pct(item['risk']['max_drawdown'])}" for item in risky) + "。")
    top = cards[0]
    lines.append(f"规则排序第一是 {top['card_name']}，动作是{top['advice']['action']}。")
    high_events = [event for event in events if event.get("impact_level") == "高"]
    if high_events:
        lines.append(f"供给端仍要看 {high_events[0]['date']} 的事件：{high_events[0]['title']}。")
    return lines[:5]


def _table_row(item: dict) -> dict:
    returns = item.get("returns") or {}
    volume = item.get("volume") or {}
    risk = item.get("risk") or {}
    return {
        "card_code": item["card_code"],
        "card_name": item["card_name"],
        "latest_price": item.get("latest_price"),
        "return_7d": (returns.get("d7") or {}).get("return"),
        "return_30d": (returns.get("d30") or {}).get("return"),
        "return_90d": (returns.get("d90") or {}).get("return"),
        "volume_7d": volume.get("d7"),
        "volume_change_7d": volume.get("change_7d"),
        "volatility": risk.get("volatility"),
        "max_drawdown": risk.get("max_drawdown"),
        "avg_active": volume.get("avg_active"),
        "highest_bid": (item.get("liquidity") or {}).get("highest_bid"),
        "divergence": (item.get("divergence") or {}).get("signal"),
        "action": item["advice"]["action"],
    }


def _divergence_table(cards):
    rows = []
    for item in cards:
        if not item.get("has_psa10"):
            continue
        rows.append(
            {
                "card_code": item["card_code"],
                "card_name": item["card_name"],
                "signal": item["divergence"]["signal"],
                "meaning": item["divergence"]["meaning"],
                "heat": item.get("heat"),
                "return_7d": (item["returns"]["d7"] or {}).get("return"),
                "volume_change_7d": (item.get("volume") or {}).get("change_7d"),
            }
        )
    return rows


def _liquidity(cards):
    rows = []
    for item in cards:
        if not item.get("has_psa10"):
            continue
        volume = item.get("volume") or {}
        liquidity = item.get("liquidity") or {}
        pop = item.get("population") or {}
        rows.append(
            {
                "card_code": item["card_code"],
                "card_name": item["card_name"],
                "active_ratio": volume.get("active_ratio"),
                "avg_active": volume.get("avg_active"),
                "median_gap_hours": volume.get("median_gap_hours"),
                "spread": liquidity.get("spread"),
                "highest_bid": liquidity.get("highest_bid"),
                "bid_count": liquidity.get("bid_count"),
                "pressure": (pop.get("pressure") or {}).get("label"),
                "pressure_score": (pop.get("pressure") or {}).get("score"),
            }
        )
    rows.sort(key=lambda row: row["avg_active"] or 0, reverse=True)
    return rows


def _forecast_row(item):
    prediction = item.get("prediction") or {}
    return {
        "card_code": item["card_code"],
        "card_name": item["card_name"],
        "available": prediction.get("available"),
        "d7": prediction.get("d7"),
        "d30": prediction.get("d30"),
        "bias": prediction.get("bias"),
        "r2": prediction.get("r2"),
        "reason": prediction.get("reason"),
        "action_hint": _forecast_hint(prediction, item["advice"]["action"]),
        "disclaimer": prediction.get("disclaimer"),
    }


def _forecast_hint(prediction, action):
    if not prediction.get("available"):
        return "样本不够，不用区间做目标。"
    bias = prediction.get("bias") or ""
    if "偏强" in bias and action not in {"卖出/减仓", "强卖出"}:
        return "统计情景偏强，已有仓位可以把止盈上移，仍按量能决定是否加仓。"
    if "偏弱" in bias:
        return "统计情景偏弱，不宜追价，等成交回到日均附近再看。"
    return "区间震荡，适合按预测高低沿做高抛低吸，而不是押单边。"


def _advice_row(item):
    advice = item["advice"]
    return {
        "card_code": item["card_code"],
        "card_name": item["card_name"],
        "action": advice["action"],
        "score": advice["score"],
        "position": advice["position"],
        "take_profit": advice["take_profit"],
        "stop_loss": advice["stop_loss"],
        "reasons": advice["reasons"],
    }


def _monitor(item):
    tips = [f"{item['card_name']}：7 日涨跌幅超过 ±20% 时系统会发站内告警。"]
    change = (item.get("volume") or {}).get("change_7d")
    if change is not None and change < -0.5:
        tips.append("近期缩量，优先看成交笔数是否回到近 30 日日均。")
    drawdown = (item.get("risk") or {}).get("max_drawdown")
    if drawdown is not None and drawdown <= -0.30:
        tips.append(f"最大回撤已到 {_pct(drawdown)}，跌破样本低点 {_cny(item.get('low'))} 时重新评估。")
    pressure = ((item.get("population") or {}).get("pressure") or {}).get("label")
    if pressure == "高":
        tips.append("供给压力高，追加生产或再抽选落地后要重算价格天花板。")
    if item.get("anomaly_count"):
        tips.append("异常成交不纳入主图，复核时单独看异常表。")
    return {"card_code": item["card_code"], "card_name": item["card_name"], "tips": tips}


def render_docx(report: dict) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), "PingFang SC")
    doc.add_heading(report["title"], level=0)
    doc.add_paragraph(f"数据截止 {report.get('as_of') or '—'} · 生成于 {report['generated_at']}")
    doc.add_paragraph(report["disclaimer"])

    doc.add_heading("1. 管理层结论", level=1)
    for line in report["executive"]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("2. 核心量化指标", level=1)
    headers = ["卡牌", "最新价", "7D", "30D", "7D量", "量变", "波动率", "最大回撤", "建议"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for row in report["metrics_table"]:
        values = [
            row["card_name"],
            _cny(row["latest_price"]),
            _pct(row["return_7d"]),
            _pct(row["return_30d"]),
            _num(row["volume_7d"]),
            _pct(row["volume_change_7d"]),
            _pct(row["volatility"]),
            _pct(row["max_drawdown"]),
            row["action"],
        ]
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value

    doc.add_heading("3. PSA10 专项分析", level=1)
    for card in report["cards"]:
        doc.add_heading(card["card_name"], level=2)
        doc.add_paragraph(card["narrative"])
        doc.add_paragraph("交互走势图在系统「报告」和「卡牌」页面，Word 导出保留文字结论。")
        if card.get("bundle_trades"):
            doc.add_paragraph(f"套出观察 {len(card['bundle_trades'])} 笔（展示最近记录，未纳入计算）。")

    doc.add_heading("4. 价格 / 成交量背离与热度", level=1)
    for row in report["divergence"]:
        doc.add_paragraph(f"{row['card_name']}：{row['signal']}。{row['meaning']} 热度 {_num(row['heat'])}。")

    doc.add_heading("5. 流动性与供需", level=1)
    for row in report["liquidity"]:
        doc.add_paragraph(
            f"{row['card_name']}：有效交易日占比 {_pct(row['active_ratio'])}，日均 { _num(row['avg_active']) } 笔，"
            f"最高竞价 {_cny(row['highest_bid'])}，供给压力 {row['pressure'] or '—'}。"
        )

    doc.add_heading("6. 异常成交", level=1)
    if not report["anomalies"]:
        doc.add_paragraph("当前没有 IQR 或业务规则命中的异常成交。")
    for block in report["anomalies"]:
        doc.add_paragraph(f"{block['card_name']} 共 {block['count']} 笔。")
        for row in block["rows"]:
            doc.add_paragraph(f"{row['date']} {row.get('time') or ''} {_cny(row['price'])}（{'；'.join(row['reasons'])}）", style="List Bullet")

    doc.add_heading("7. 未来价格预测", level=1)
    for row in report["forecasts"]:
        if not row["available"]:
            doc.add_paragraph(f"{row['card_name']}：{row['reason']}。")
            continue
        doc.add_paragraph(
            f"{row['card_name']}：7 日区间 {_cny(row['d7']['low'])}–{_cny(row['d7']['high'])}，"
            f"30 日区间 {_cny(row['d30']['low'])}–{_cny(row['d30']['high'])}，基准判断 {row['bias']}。{row['action_hint']}"
        )
    if report["forecasts"]:
        doc.add_paragraph(report["forecasts"][0]["disclaimer"])

    doc.add_heading("8. 投资建议", level=1)
    for row in report["advice"]["ranking"]:
        doc.add_paragraph(
            f"{row['card_name']} · {row['action']}（评分 {row['score']}）。仓位：{row['position']}。"
            f"止盈 {_cny(row['take_profit'])}，止损 {_cny(row['stop_loss'])}。"
        )
        for reason in row["reasons"]:
            doc.add_paragraph(reason, style="List Bullet")
    doc.add_paragraph("风险控制")
    for line in report["advice"]["risk_controls"]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("9. 外部市场核验与事件", level=1)
    doc.add_paragraph(report["external"]["note"])
    for row in report["external"]["population"]:
        doc.add_paragraph(f"{row['card_name']}：总人口 {row['total'] or '—'}，PSA10 {row['psa10'] or '—'}，Gem Rate {_pct(row['gem_rate'])}，供给压力 {row['pressure'] or '—'}。")
    for event in report["external"]["events"]:
        doc.add_paragraph(f"{event['date']} [{event['impact_level'] or '—'}] {event['title']}。{event.get('description') or ''}")

    doc.add_heading("10. 建议后续监控", level=1)
    for block in report["monitoring"]:
        for tip in block["tips"]:
            doc.add_paragraph(tip, style="List Bullet")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_xlsx(report: dict) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "核心指标"
    headers = ["代码", "名称", "最新价", "7D收益", "30D收益", "90D收益", "7D量", "7D量变", "波动率", "最大回撤", "日均量", "最高竞价", "背离", "建议"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in report["metrics_table"]:
        sheet.append([
            row["card_code"], row["card_name"], row["latest_price"], row["return_7d"], row["return_30d"], row["return_90d"],
            row["volume_7d"], row["volume_change_7d"], row["volatility"], row["max_drawdown"], row["avg_active"],
            row["highest_bid"], row["divergence"], row["action"],
        ])
    events = book.create_sheet("事件")
    events.append(["日期", "类型", "卡牌", "影响", "标题", "说明"])
    for event in report["external"]["events"]:
        events.append([event["date"], event["type"], event["card_code"], event["impact_level"], event["title"], event["description"]])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _pct(value) -> str:
    if value is None:
        return "N/A"
    return f"{value * 100:+.1f}%"


def _cny(value) -> str:
    if value is None:
        return "—"
    number = float(value)
    if abs(number - round(number)) < 0.001:
        return f"¥{number:,.0f}"
    return f"¥{number:,.2f}"


def _num(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value)
